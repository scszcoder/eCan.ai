"""商品上架 spreadsheet -> listing.json folders, for the 商品上架 skills.

A store owner fills one spreadsheet (one row per product, per SKU, per picture,
per size) and drops it with the pictures into the store's 上架 folder; each
product becomes ``<out>/<商品编码>/listing.json`` plus its files, which is what
the 上架 skills publish. Customers keep spreadsheets, not JSON.

Sheets (``write_template`` writes an empty one with instructions):
  商品   商品编码*  商品标题*  类目*  品牌  价格*  库存  属性  描述
  规格   商品编码*  颜色  尺码  价格  库存  规格编码
  图片   商品编码*  文件名*  用途*（主图/详情图/尺码表图/颜色图/视频）  颜色
  尺码表 商品编码*  尺码*  then any measurement columns (身高cm, 体重斤, 衣长cm …)
"""
from __future__ import annotations

import json
import os
import re
import shutil
from typing import Any

SHEETS = {
    "商品": ["商品编码", "商品标题", "类目", "品牌", "价格", "库存", "属性", "描述"],
    "规格": ["商品编码", "颜色", "尺码", "价格", "库存", "规格编码"],
    "图片": ["商品编码", "文件名", "用途", "颜色"],
    "尺码表": ["商品编码", "尺码", "身高cm", "体重斤", "衣长cm", "胸围cm"],
}
HELP = [
    ("商品编码", "必填。同一商品在各表里用同一个编码；也是这个商品的文件夹名。"),
    ("商品标题", "必填。按平台要求的标题原样填写。"),
    ("类目", "必填。从大到小用 / 分隔，例如：童装/婴儿装/亲子装/T恤。"),
    ("价格 / 库存", "商品表里填默认值；规格表里某个 SKU 的价格或库存不同才填。"),
    ("属性", "多条用分号分隔，每条“名称=值”，例如：面料=棉;领型=圆领;适用季节=夏季。"),
    ("规格", "每个 SKU 一行（颜色 × 尺码）。没有颜色或尺码的商品可以只填一列。"),
    ("图片", "文件名是和表格放在同一个文件夹里的图片/视频文件。用途：主图（按顺序，第一张是首图）、详情图、尺码表图、颜色图（同时填颜色）、视频。"),
    ("尺码表", "每个尺码一行；尺寸列可以增删（列名写清单位，如 身高cm、体重斤）。没有尺码表就不填，平台要求时会报告缺少。"),
]
ROLE = {"主图": "main", "详情图": "detail", "尺码表图": "size_chart", "颜色图": "color", "视频": "video"}


def write_template(path: str) -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        wb.create_sheet(name).append(cols)
    ws = wb.create_sheet("填写说明")
    ws.append(["项目", "说明"])
    for row in HELP:
        ws.append(list(row))
    wb.save(path)
    return path


def _num(v: Any) -> Any:
    if v is None or str(v).strip() == "":
        return None
    try:
        f = float(str(v).strip())
        return int(f) if f.is_integer() else f
    except ValueError:
        return str(v).strip()


def _rows(wb, name: str) -> list[dict]:
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    head = [str(h).strip() if h is not None else "" for h in next(it, [])]
    out = []
    for r in it:
        d = {h: (str(v).strip() if isinstance(v, str) else v) for h, v in zip(head, r) if h}
        if any(v not in (None, "") for v in d.values()):
            out.append(d)
    return out


def _attrs(text: Any) -> dict:
    out = {}
    for part in re.split(r"[;；\n]", str(text or "")):
        if "=" in part or "＝" in part:
            k, v = re.split(r"[=＝]", part, maxsplit=1)
            if k.strip():
                out[k.strip()] = v.strip()
    return out


def read_sheet(xlsx: str) -> dict[str, dict]:
    """The listings in a 商品上架 spreadsheet, by 商品编码 (file names not resolved)."""
    import io
    import openpyxl
    # From bytes: openpyxl keeps a file it opened locked on Windows, and the
    # owner could not save or replace their spreadsheet while the app ran.
    with open(xlsx, "rb") as f:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    out: dict[str, dict] = {}
    for r in _rows(wb, "商品"):
        code = str(r.get("商品编码") or "").strip()
        if not code:
            continue
        cat = [c.strip() for c in re.split(r"[/>／＞]", str(r.get("类目") or "")) if c.strip()]
        attrs = _attrs(r.get("属性"))
        if r.get("品牌"):
            attrs.setdefault("品牌", str(r["品牌"]))
        out[code] = {"code": code, "title": str(r.get("商品标题") or ""), "category": cat, "attributes": attrs,
                     "price": _num(r.get("价格")), "stock": _num(r.get("库存")),
                     "description": str(r.get("描述") or ""), "skus": [], "files": {}}
    for r in _rows(wb, "规格"):
        L = out.get(str(r.get("商品编码") or "").strip())
        if L is None:
            continue
        sku = {k: str(r[k]) for k in ("颜色", "尺码") if r.get(k) not in (None, "")}
        sku["价格"] = _num(r.get("价格")) if r.get("价格") not in (None, "") else L["price"]
        sku["库存"] = _num(r.get("库存")) if r.get("库存") not in (None, "") else L["stock"]
        if r.get("规格编码"):
            sku["商家编码"] = str(r["规格编码"])
        L["skus"].append(sku)
    for r in _rows(wb, "图片"):
        L = out.get(str(r.get("商品编码") or "").strip())
        role = ROLE.get(str(r.get("用途") or "").strip())
        name = str(r.get("文件名") or "").strip()
        if L is None or not role or not name:
            continue
        if role == "color":
            L["files"].setdefault("color", {})[str(r.get("颜色") or "")] = name
        elif role in ("main", "detail"):
            L["files"].setdefault(role, []).append(name)
        else:
            L["files"][role] = name
    for r in _rows(wb, "尺码表"):
        L = out.get(str(r.get("商品编码") or "").strip())
        if L is None or r.get("尺码") in (None, ""):
            continue
        row = {k: str(_num(v)) for k, v in r.items() if k != "商品编码" and v not in (None, "")}
        chart = L.setdefault("size_chart", {"columns": [], "rows": []})
        for k in row:
            if k not in chart["columns"]:
                chart["columns"].append(k)
        chart["rows"].append(row)
    for L in out.values():
        for spec, key in (("颜色", "colors"), ("尺码", "sizes")):
            vals = list(dict.fromkeys(s[spec] for s in L["skus"] if s.get(spec)))
            if vals:
                L[key] = vals
    return out


def problems(listing: dict, folder: str = "") -> list[str]:
    """What stops this listing from being published, in words for the store owner."""
    out = []
    for key, what in (("title", "商品标题"), ("category", "类目"), ("price", "价格")):
        if not listing.get(key):
            out.append(f"缺少{what}")
    if not (listing.get("files") or {}).get("main"):
        out.append("没有主图（图片表里用途填“主图”）")
    for s in listing.get("skus") or []:
        if s.get("价格") in (None, "") or s.get("库存") in (None, ""):
            out.append(f"规格 {s.get('颜色', '')} {s.get('尺码', '')} 缺少价格或库存".replace("  ", " "))
    if folder:
        for v in (listing.get("files") or {}).values():
            for name in (v.values() if isinstance(v, dict) else v if isinstance(v, list) else [v]):
                if name and not os.path.isfile(os.path.join(folder, name)):
                    out.append(f"找不到文件 {name}")
    return out


def convert(xlsx: str, out_root: str, media_dir: str = "") -> list[tuple[str, str, list[str]]]:
    """Write ``<out_root>/<商品编码>/listing.json`` (+ its files, copied from
    ``media_dir``, default: the spreadsheet's folder) for every product.
    Returns ``(code, folder, problems)``; a product with problems is still
    written, so the owner can see and fix it, but the skill skips it."""
    media_dir = media_dir or os.path.dirname(os.path.abspath(xlsx))
    done = []
    for code, L in read_sheet(xlsx).items():
        folder = os.path.join(out_root, re.sub(r'[\\/:*?"<>|]', "-", code))
        os.makedirs(folder, exist_ok=True)
        for v in L["files"].values():
            for name in (v.values() if isinstance(v, dict) else v if isinstance(v, list) else [v]):
                src = os.path.join(media_dir, name)
                if name and os.path.isfile(src) and not os.path.isfile(os.path.join(folder, name)):
                    shutil.copyfile(src, os.path.join(folder, name))
        found = problems(L, folder)
        if found:
            L["problems"] = found
        with open(os.path.join(folder, "listing.json"), "w", encoding="utf-8") as f:
            json.dump(L, f, ensure_ascii=False, indent=1)
        done.append((code, folder, found))
    return done
