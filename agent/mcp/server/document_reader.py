"""Read office / PDF / text documents as plain text for an LLM (the os_read_document tool).

xlsx/xlsm -> one block per sheet, rows as tab-separated lines; csv/tsv -> rows;
pdf -> text per page (PyMuPDF, PyPDF2 as fallback); docx -> paragraphs + tables;
pptx -> text per slide + notes; anything else -> decoded as text.
Legacy binary formats (.xls/.doc/.ppt) are reported as unsupported.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Dict, List, Optional

TEXT_ENCODINGS = ("utf-8-sig", "gb18030", "latin-1")
LEGACY = {".xls": "xlsx", ".doc": "docx", ".ppt": "pptx"}


def _decode(raw: bytes) -> str:
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\t", " ").replace("\r", " ").replace("\n", " ").strip()


def _read_xlsx(path: Path, max_rows: int, sheet: Optional[str]) -> Dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        picked = [sheet] if sheet else names
        blocks: List[str] = []
        truncated = False
        for name in picked:
            if name not in names:
                blocks.append(f"## Sheet: {name}\n(no such sheet; sheets: {', '.join(names)})")
                continue
            lines = []
            for row in wb[name].iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                if len(lines) >= max_rows:
                    truncated = True
                    lines.append(f"... (more rows; showing the first {max_rows})")
                    break
                lines.append("\t".join(cells))
            blocks.append(f"## Sheet: {name}\n" + ("\n".join(lines) if lines else "(empty)"))
        return {"type": "xlsx", "sheets": names, "content": "\n\n".join(blocks), "truncated": truncated}
    finally:
        wb.close()


def _read_csv(path: Path, max_rows: int) -> Dict:
    text = _decode(path.read_bytes())
    delim = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        delim = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|").delimiter
    except csv.Error:
        pass
    lines, truncated = [], False
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        if not any(c.strip() for c in row):
            continue
        if len(lines) >= max_rows:
            truncated = True
            lines.append(f"... (more rows; showing the first {max_rows})")
            break
        lines.append("\t".join(_cell(c) for c in row))
    return {"type": "csv", "content": "\n".join(lines), "truncated": truncated}


def _read_pdf(path: Path) -> Dict:
    pages: List[str] = []
    try:
        import fitz  # PyMuPDF
        with fitz.open(path) as doc:
            pages = [p.get_text() for p in doc]
    except Exception:
        from PyPDF2 import PdfReader
        pages = [(p.extract_text() or "") for p in PdfReader(str(path)).pages]
    out = {"type": "pdf", "pages": len(pages),
           "content": "\n\n".join(f"## Page {i + 1}\n{t.strip()}" for i, t in enumerate(pages))}
    if pages and not any(t.strip() for t in pages):
        out["note"] = "no text layer (scanned PDF?) - the pages are images; OCR is needed to read them"
    return out


def _read_docx(path: Path) -> Dict:
    import docx
    d = docx.Document(str(path))
    parts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for ti, table in enumerate(d.tables, 1):
        rows = ["\t".join(_cell(c.text) for c in r.cells) for r in table.rows]
        parts.append(f"## Table {ti}\n" + "\n".join(rows))
    return {"type": "docx", "content": "\n".join(parts)}


def _read_pptx(path: Path) -> Dict:
    from pptx import Presentation
    prs = Presentation(str(path))
    blocks = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                texts.extend("\t".join(_cell(c.text) for c in r.cells) for r in shape.table.rows)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        blocks.append(f"## Slide {i}\n" + "\n".join(texts))
    return {"type": "pptx", "slides": len(prs.slides), "content": "\n\n".join(blocks)}


def read_document(file_path: str, max_chars: int = 40000, max_rows: int = 300,
                  sheet: Optional[str] = None) -> Dict:
    path = Path(file_path)
    if not path.is_file():
        return {"file_path": str(path), "error": "file not found"}
    ext = path.suffix.lower()
    base = {"file_path": str(path), "size_bytes": path.stat().st_size}
    if ext in LEGACY:
        return {**base, "error": f"legacy {ext} format is not supported; save it as .{LEGACY[ext]} and read that"}
    if ext in (".xlsx", ".xlsm"):
        out = _read_xlsx(path, max_rows, sheet)
    elif ext in (".csv", ".tsv"):
        out = _read_csv(path, max_rows)
    elif ext == ".pdf":
        out = _read_pdf(path)
    elif ext == ".docx":
        out = _read_docx(path)
    elif ext == ".pptx":
        out = _read_pptx(path)
    else:
        out = {"type": "text", "content": _decode(path.read_bytes())}
    content = out.get("content") or ""
    if len(content) > max_chars:
        out["content"] = content[:max_chars] + f"\n... (cut at {max_chars} of {len(content)} characters)"
        out["truncated"] = True
    out.setdefault("truncated", False)
    return {**base, **out}
