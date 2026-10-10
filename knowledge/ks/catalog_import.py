"""
Catalog import: CSV / XLSX / JSON -> canonical product records.

Pipeline (each step is pure and unit-tested):
  read_table()       file -> headers + rows (encoding detected)
  suggest_mapping()  merchant headers -> canonical fields (zh/en synonyms)
  build_products()   rows + mapping -> cleaned product dicts, variants merged,
                     plus a validation report (errors, warnings, skipped rows)

Price and stock columns are recognised but NOT stored: they change too often
and are read live from the product card / platform tools instead.
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Tuple

from knowledge.ks.text import clean_text

# Canonical fields. "attribute" is the catch-all for anything product-specific.
FIELDS = (
    "code", "title", "platform_item_id", "store", "category", "brand",
    "description", "selling_points", "care", "faq_q", "faq_a",
    "variant_sku", "color", "size", "material",
    "price", "stock", "image", "attribute", "ignore",
)

FIELD_LABELS = {
    "code": "商品编码 / Product code", "title": "商品标题 / Title",
    "platform_item_id": "平台商品ID / Platform item ID", "store": "店铺 / Store",
    "category": "类目 / Category", "brand": "品牌 / Brand",
    "description": "描述 / Description", "selling_points": "卖点 / Selling points",
    "care": "洗涤保养 / Care", "faq_q": "常见问题 / FAQ question", "faq_a": "回答 / FAQ answer",
    "variant_sku": "SKU编码 / Variant SKU", "color": "颜色 / Color", "size": "尺码 / Size",
    "material": "材质 / Material", "price": "价格（不保存）/ Price (not stored)",
    "stock": "库存（不保存）/ Stock (not stored)", "image": "图片 / Image",
    "attribute": "商品属性 / Attribute", "ignore": "忽略 / Ignore",
}

# Header synonyms, matched after normalisation (lowercase, no spaces/punctuation).
SYNONYMS: Dict[str, List[str]] = {
    "code": ["商品编码", "商家编码", "货号", "款号", "商品货号", "spu", "spu编码", "productcode", "product_code", "itemcode", "code", "outerid", "outer_id"],
    "title": ["商品名称", "商品标题", "宝贝标题", "宝贝名称", "标题", "名称", "品名", "title", "name", "productname"],
    "platform_item_id": ["商品id", "宝贝id", "平台商品id", "itemid", "item_id", "numiid", "num_iid", "goodsid", "goods_id"],
    "store": ["店铺", "店铺名称", "店铺名", "store", "shop", "storename", "shopname"],
    "category": ["类目", "分类", "商品类目", "叶子类目", "category", "categoryname"],
    "brand": ["品牌", "brand"],
    "description": ["描述", "商品描述", "详情", "商品详情", "宝贝描述", "description", "desc", "detail"],
    "selling_points": ["卖点", "商品卖点", "亮点", "sellingpoints", "highlights", "features"],
    "care": ["洗涤说明", "洗涤", "保养", "洗护", "care", "careinstructions", "washing"],
    "faq_q": ["问题", "常见问题", "faq", "question", "q"],
    "faq_a": ["回答", "答案", "answer", "a"],
    "variant_sku": ["sku编码", "sku", "skuid", "sku_id", "规格编码", "variantsku"],
    "color": ["颜色", "颜色分类", "颜色名称", "color", "colour"],
    "size": ["尺码", "尺寸", "码数", "size"],
    "material": ["材质", "面料", "成分", "材料", "面料成分", "material", "fabric", "composition"],
    "price": ["价格", "售价", "一口价", "单价", "price", "saleprice"],
    "stock": ["库存", "数量", "可售数量", "stock", "inventory", "qty", "quantity"],
    "image": ["图片", "主图", "图片链接", "image", "imageurl", "picture"],
}

MAX_ROWS = 50000
ROW_KEY = "__row__"  # original line/row number, so reports point at the merchant's file


@dataclass
class ImportReport:
    rows_total: int = 0
    products: int = 0
    variants: int = 0
    listings: int = 0
    skipped: List[Dict[str, Any]] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    ignored_columns: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "rows_total": self.rows_total, "products": self.products, "variants": self.variants,
            "listings": self.listings, "skipped": self.skipped[:200], "skipped_count": len(self.skipped),
            "warnings": self.warnings[:200], "warning_count": len(self.warnings),
            "errors": self.errors, "ignored_columns": self.ignored_columns,
        }


def _norm_header(h: str) -> str:
    return re.sub(r"[\s_\-/（）()·:：]+", "", str(h or "").strip().lower())


def suggest_mapping(headers: Iterable[str]) -> Dict[str, str]:
    """Map each header to a canonical field; unknown headers become attributes."""
    mapping: Dict[str, str] = {}
    used: set = set()
    norm_syn = {f: {_norm_header(s) for s in syns} for f, syns in SYNONYMS.items()}
    for h in headers:
        nh = _norm_header(h)
        field_name = None
        for f, syns in norm_syn.items():
            if nh in syns:
                field_name = f
                break
        if field_name is None:
            # Partial match for long headers like "面料成分(%)".
            for f, syns in norm_syn.items():
                if any(len(s) >= 2 and s in nh for s in syns) and f not in ("faq_a", "faq_q"):
                    field_name = f
                    break
        if field_name in used and field_name not in ("attribute", "image"):
            field_name = "attribute"  # never map two columns to one single-value field
        mapping[h] = field_name or ("attribute" if nh else "ignore")
        used.add(mapping[h])
    return mapping


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    try:
        import chardet
        enc = chardet.detect(data[:200000]).get("encoding") or "utf-8"
        return data.decode(enc, errors="replace")
    except Exception:
        return data.decode("utf-8", errors="replace")


def read_table(path: str) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Read a CSV / XLSX / JSON catalog into (headers, rows)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        with open(path, "rb") as f:
            text = _decode(f.read())
        dialect = csv.excel_tab if ext == ".tsv" else csv.excel
        if ext != ".tsv":
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
            except csv.Error:
                pass
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = [r for r in reader]
        if not all_rows:
            return [], []
        headers = [h.strip() for h in all_rows[0]]
        rows = [{**dict(zip(headers, r)), ROW_KEY: n} for n, r in enumerate(all_rows[1:MAX_ROWS + 1], start=2)
                if any(c.strip() for c in r)]
        return headers, rows
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            return [], []
        headers = [str(h).strip() if h is not None else f"列{i + 1}" for i, h in enumerate(header_row)]
        rows = []
        for n, r in enumerate(it, start=2):
            if len(rows) >= MAX_ROWS:
                break
            if r is None or not any(v not in (None, "") for v in r):
                continue
            rows.append({**{h: ("" if v is None else v) for h, v in zip(headers, r)}, ROW_KEY: n})
        wb.close()
        return headers, rows
    if ext == ".json":
        with open(path, "rb") as f:
            data = json.loads(_decode(f.read()))
        if isinstance(data, dict):
            data = data.get("products") or data.get("items") or data.get("data") or []
        if not isinstance(data, list):
            raise ValueError("JSON catalog must be a list of objects or {\"products\": [...]}")
        rows = [{**r, ROW_KEY: n} for n, r in enumerate(data[:MAX_ROWS], start=1) if isinstance(r, dict)]
        headers: List[str] = []
        for r in rows:
            for k in r.keys():
                if k != ROW_KEY and k not in headers:
                    headers.append(k)
        return headers, rows
    raise ValueError(f"Unsupported catalog file type: {ext} (use .csv, .xlsx or .json)")


def _as_list(value: Any) -> List[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [clean_text(v) for v in value if clean_text(v)]
    parts = re.split(r"[\n;；|]+", clean_text(value))
    return [p.strip(" -•·") for p in parts if p.strip(" -•·")]


def _cell(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel turns 123456 into 123456.0
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return clean_text(value)


def build_products(rows: List[Dict[str, Any]], mapping: Dict[str, str]) -> Tuple[List[Dict[str, Any]], ImportReport]:
    """Turn mapped rows into canonical product dicts. Rows sharing a product
    code (or, without codes, a title) are merged and their color/size become
    variants."""
    report = ImportReport(rows_total=len(rows))
    report.ignored_columns = [h for h, f in mapping.items() if f in ("price", "stock")]
    if report.ignored_columns:
        report.warnings.append(
            "价格/库存列不会保存（会从商品卡片或平台实时读取）: " + ", ".join(report.ignored_columns))
    has_code = "code" in mapping.values()
    if "title" not in mapping.values():
        report.errors.append("缺少商品标题列 / A title column is required")
        return [], report

    products: Dict[str, Dict[str, Any]] = {}
    for i, row in enumerate(rows, start=2):  # row 1 is the header
        i = row.get(ROW_KEY, i)  # the row number in the merchant's file
        rec: Dict[str, Any] = {"attributes": {}, "listings": []}
        multi: Dict[str, List[str]] = {}
        for header, field_name in mapping.items():
            raw = row.get(header)
            val = _cell(raw)
            if field_name in ("ignore", "price", "stock") or val == "":
                continue
            if field_name == "attribute":
                rec["attributes"][str(header).strip()] = val
            elif field_name in ("selling_points", "image"):
                multi.setdefault(field_name, []).extend(_as_list(raw))
            else:
                rec[field_name] = val
        title = rec.get("title", "")
        if not title:
            report.skipped.append({"row": i, "reason": "缺少标题 / missing title"})
            continue
        code = rec.get("code") or ("" if has_code else title)
        if not code:
            report.skipped.append({"row": i, "reason": "缺少商品编码 / missing product code"})
            continue
        key = code.lower()
        p = products.get(key)
        if p is None:
            p = {
                "code": code, "title": title, "category": rec.get("category", ""),
                "brand": rec.get("brand", ""), "description": rec.get("description", ""),
                "care": rec.get("care", ""), "attributes": {}, "variants": [],
                "selling_points": [], "images": [], "faq": [], "listings": [], "_rows": [],
            }
            products[key] = p
        p["_rows"].append(i)
        if title != p["title"]:
            report.warnings.append(f"第{i}行: 商品 {code} 的标题与前面不同，保留第一个 / title differs, kept the first")
        if rec.get("material"):
            prev = p["attributes"].get("材质")
            if prev and prev != rec["material"]:
                report.warnings.append(f"第{i}行: 商品 {code} 材质冲突「{prev}」vs「{rec['material']}」/ material conflict")
            else:
                p["attributes"]["材质"] = rec["material"]
        for k, v in rec["attributes"].items():
            prev = p["attributes"].get(k)
            if prev and prev != v and k not in ("颜色", "尺码"):
                report.warnings.append(f"第{i}行: 商品 {code} 属性「{k}」冲突「{prev}」vs「{v}」")
            else:
                p["attributes"].setdefault(k, v)
        for f in ("category", "brand", "description", "care"):
            if not p[f] and rec.get(f):
                p[f] = rec[f]
        variant = {k: rec[k] for k in ("variant_sku", "color", "size") if rec.get(k)}
        if variant:
            v = {"sku": variant.get("variant_sku", ""), "颜色": variant.get("color", ""), "尺码": variant.get("size", "")}
            v = {k: x for k, x in v.items() if x}
            if v not in p["variants"]:
                p["variants"].append(v)
        for sp in multi.get("selling_points", []):
            if sp not in p["selling_points"]:
                p["selling_points"].append(sp)
        for img in multi.get("image", []):
            if img not in p["images"]:
                p["images"].append(img)
        if rec.get("faq_q") and rec.get("faq_a"):
            qa = {"q": rec["faq_q"], "a": rec["faq_a"]}
            if qa not in p["faq"]:
                p["faq"].append(qa)
        if rec.get("platform_item_id"):
            lst = {"platform_item_id": rec["platform_item_id"], "store": rec.get("store", "")}
            if lst not in p["listings"]:
                p["listings"].append(lst)

    out = list(products.values())
    for p in out:
        p.pop("_rows", None)
        # Variant colors/sizes summarised as attributes so they're searchable.
        colors = sorted({v["颜色"] for v in p["variants"] if v.get("颜色")})
        sizes = [s for s in dict.fromkeys(v["尺码"] for v in p["variants"] if v.get("尺码"))]
        if colors:
            p["attributes"].setdefault("可选颜色", "、".join(colors))
        if sizes:
            p["attributes"].setdefault("可选尺码", "、".join(sizes))
    report.products = len(out)
    report.variants = sum(len(p["variants"]) for p in out)
    report.listings = sum(len(p["listings"]) for p in out)
    return out, report


def preview_rows(rows: List[Dict[str, Any]], n: int = 5) -> List[Dict[str, str]]:
    return [{k: _cell(v)[:120] for k, v in r.items() if k != ROW_KEY} for r in rows[:n]]
