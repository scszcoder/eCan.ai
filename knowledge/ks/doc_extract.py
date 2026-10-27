"""
Single-document ingestion: extract text/tables from a product document and
draft a product record for the merchant to confirm.

The draft is heuristic (key: value lines -> attributes, Q/A pairs -> FAQ,
tables -> size chart, bullet lines -> selling points). An LLM extractor can
replace ``draft_product`` later behind the same return shape; the merchant
always reviews the draft before it is saved.
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Tuple

from knowledge.ks.text import attribute_family, clean_text

SUPPORTED_DOC_EXTS = (".pdf", ".docx", ".txt", ".md", ".markdown")

_KV_RE = re.compile(r"^\s*([^:：\n]{1,16})\s*[:：]\s*(.{1,300})$")
_Q_RE = re.compile(r"^\s*(?:Q|问|问题)\s*[0-9]*\s*[:：.、]\s*(.+)$", re.I)
_A_RE = re.compile(r"^\s*(?:A|答|回答)\s*[0-9]*\s*[:：.、]\s*(.+)$", re.I)
_BULLET_RE = re.compile(r"^\s*(?:[-•·●▪✓✔]|\d+[.、)])\s*(.{4,200})$")
_SIZE_HEADER_HINTS = ("尺码", "码数", "size", "胸围", "衣长", "腰围", "身高", "肩宽", "袖长", "裤长", "臀围")


def extract_document(path: str) -> Tuple[str, List[List[List[str]]]]:
    """Return (text, tables). Tables are lists of rows of cells."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_DOC_EXTS:
        raise ValueError(f"Unsupported document type {ext}; supported: {', '.join(SUPPORTED_DOC_EXTS)}")
    tables: List[List[List[str]]] = []
    if ext == ".pdf":
        import fitz  # PyMuPDF
        parts = []
        with fitz.open(path) as doc:
            for page in doc:
                parts.append(page.get_text("text"))
                try:
                    for t in page.find_tables().tables:
                        rows = [[clean_text(c) for c in row] for row in t.extract()]
                        if rows:
                            tables.append(rows)
                except Exception:
                    pass  # table detection is best-effort
        text = "\n".join(parts)
    elif ext == ".docx":
        import docx
        d = docx.Document(path)
        text = "\n".join(p.text for p in d.paragraphs)
        for t in d.tables:
            rows = [[clean_text(c.text) for c in row.cells] for row in t.rows]
            if rows:
                tables.append(rows)
    else:
        with open(path, "rb") as f:
            raw = f.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030", errors="replace")
        tables.extend(_markdown_tables(text))
    return clean_text(text), tables


def _markdown_tables(text: str) -> List[List[List[str]]]:
    tables, cur = [], []
    for line in text.splitlines():
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # separator row
            cur.append(cells)
        elif cur:
            tables.append(cur)
            cur = []
    if cur:
        tables.append(cur)
    return tables


# Any category: a size column (尺码 / 鞋码 / 规格 / size …) next to a measurement column
# (a cm/mm/kg unit, or a length/girth word: 脚长, 头围, 腰围 …) -- not a clothing word list.
_SIZE_LABEL_RE = re.compile(r"尺码|码数|鞋码|码号|号型|规格|型号|圈号|size", re.I)
_MEASURE_RE = re.compile(r"cm|mm|厘米|毫米|公分|kg|公斤|斤|长|围|宽|高|深|身高|体重", re.I)


def _is_size_table(table: List[List[str]]) -> bool:
    if len(table) < 2:
        return False
    head = [str(h) for h in table[0]]
    joined = " ".join(head).lower()
    if sum(1 for h in _SIZE_HEADER_HINTS if h in joined) >= 2 or ("尺码" in joined or "size" in joined):
        return True
    label = [i for i, h in enumerate(head) if _SIZE_LABEL_RE.search(h)]
    return bool(label) and any(_MEASURE_RE.search(h) for i, h in enumerate(head) if i not in label)


def draft_product(path: str, text: str, tables: List[List[List[str]]]) -> Dict[str, Any]:
    """Heuristic product draft from document text. Every field is editable."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    filename_title = os.path.splitext(os.path.basename(path))[0]
    attributes: Dict[str, str] = {}
    faq: List[Dict[str, str]] = []
    selling_points: List[str] = []
    care_lines: List[str] = []
    body: List[str] = []
    pending_q = None
    title = ""
    code = ""
    category = ""
    brand = ""

    for line in lines:
        qm, am = _Q_RE.match(line), _A_RE.match(line)
        if qm:
            pending_q = qm.group(1).strip()
            continue
        if am and pending_q:
            faq.append({"q": pending_q, "a": am.group(1).strip()})
            pending_q = None
            continue
        kv = _KV_RE.match(line)
        if kv:
            key, val = kv.group(1).strip(), kv.group(2).strip()
            lk = key.lower()
            if lk in ("商品名称", "产品名称", "品名", "名称", "标题", "title", "name", "product"):
                title = title or val
            elif lk in ("货号", "款号", "商品编码", "型号", "商家编码", "code", "model", "sku"):
                code = code or val
            elif lk in ("类目", "分类", "品类", "category"):
                category = category or val
            elif lk in ("品牌", "brand"):
                brand = brand or val
            elif attribute_family(key) == "care":
                care_lines.append(val)
            else:
                attributes.setdefault(key, val)
            continue
        bm = _BULLET_RE.match(line)
        if bm and len(selling_points) < 12:
            selling_points.append(bm.group(1).strip())
            continue
        body.append(line)

    size_chart: Dict[str, Any] = {}
    other_tables = []
    for t in tables:
        if not size_chart and _is_size_table(t):
            size_chart = {"headers": t[0], "rows": t[1:]}
        else:
            other_tables.append(t)
    # Two-column tables are usually attribute sheets.
    for t in other_tables:
        if all(len(r) == 2 for r in t):
            for k, v in t:
                if k and v and k not in attributes:
                    attributes[k] = v

    if not title:
        title = next((l for l in lines if 4 <= len(l) <= 60 and not _KV_RE.match(l)), filename_title)
    return {
        "code": code or filename_title,
        "title": title,
        "category": category,
        "brand": brand,
        "attributes": attributes,
        "variants": [],
        "size_chart": size_chart,
        "selling_points": selling_points,
        "description": "\n".join(body)[:20000],
        "care": "；".join(care_lines),
        "faq": faq,
        "images": [],
    }
