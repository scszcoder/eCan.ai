"""Catalog import, cleaning report, single-document drafts, jobs and re-index."""
import csv
import json

import pytest

from knowledge.ks import catalog_import as ci
from knowledge.ks import doc_extract
from knowledge.ks.embedding import HashingEmbedder

HEADERS = ["商家编码", "宝贝标题", "宝贝ID", "店铺名称", "类目", "颜色", "尺码", "面料成分", "一口价", "库存", "卖点", "领型"]
ROWS = [
    ["TS-100", "女士<b>纯棉</b>短袖", "7700001", "小铺天猫店", "女装/T恤", "白", "S", "100%棉", "59", "100", "透气;不起球", "圆领"],
    ["TS-100", "女士<b>纯棉</b>短袖", "7700001", "小铺天猫店", "女装/T恤", "黑", "M", "100%棉", "59", "80", "透气", "圆领"],
    ["TS-100", "女士纯棉短袖", "4400001", "小铺抖音店", "女装/T恤", "黑", "L", "95%棉 5%氨纶", "59", "10", "", "圆领"],
    ["", "", "", "", "", "", "", "", "", "", "", ""],
    ["TS-200", "", "7700002", "小铺天猫店", "女装/T恤", "", "", "", "", "", "", ""],
    ["TS-300", "女士亚麻衬衫", "7700003", "小铺天猫店", "女装/衬衫", "米", "M", "100%亚麻", "129", "5", "", "翻领"],
]


def _write_csv(path, encoding="gbk"):
    with open(path, "w", encoding=encoding, newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(ROWS)


def test_mapping_recognises_platform_export_headers():
    m = ci.suggest_mapping(HEADERS)
    assert m["商家编码"] == "code" and m["宝贝标题"] == "title" and m["宝贝ID"] == "platform_item_id"
    assert m["店铺名称"] == "store" and m["面料成分"] == "material"
    assert m["一口价"] == "price" and m["库存"] == "stock"
    assert m["领型"] == "attribute"


def test_gbk_csv_is_decoded(tmp_path):
    p = tmp_path / "catalog.csv"
    _write_csv(p, "gbk")
    headers, rows = ci.read_table(str(p))
    assert headers[1] == "宝贝标题"
    assert len(rows) == 5  # the blank row is dropped


def test_build_products_merges_variants_and_reports(tmp_path):
    p = tmp_path / "catalog.csv"
    _write_csv(p, "utf-8-sig")
    headers, rows = ci.read_table(str(p))
    products, report = ci.build_products(rows, ci.suggest_mapping(headers))
    by_code = {x["code"]: x for x in products}
    assert set(by_code) == {"TS-100", "TS-300"}
    shirt = by_code["TS-100"]
    assert shirt["title"] == "女士纯棉短袖"  # HTML stripped
    assert len(shirt["variants"]) == 3
    assert shirt["attributes"]["可选尺码"] == "S、M、L"
    assert shirt["selling_points"] == ["透气", "不起球"]
    assert {l["platform_item_id"] for l in shirt["listings"]} == {"7700001", "4400001"}
    assert "一口价" not in shirt["attributes"] and "库存" not in shirt["attributes"]
    assert any("材质冲突" in w for w in report.warnings)
    assert report.skipped == [{"row": 6, "reason": "缺少标题 / missing title"}]
    assert any("价格/库存" in w for w in report.warnings)


def test_missing_title_column_is_an_error():
    _, report = ci.build_products([{"x": "1"}], {"x": "attribute"})
    assert report.errors


def test_xlsx_and_json_catalogs(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["货号", "商品名称", "材质"])
    ws.append([12345, "儿童雨靴", "PVC"])  # numeric code must not become 12345.0
    xp = tmp_path / "c.xlsx"
    wb.save(xp)
    headers, rows = ci.read_table(str(xp))
    products, _ = ci.build_products(rows, ci.suggest_mapping(headers))
    assert products[0]["code"] == "12345" and products[0]["attributes"]["材质"] == "PVC"

    jp = tmp_path / "c.json"
    jp.write_text(json.dumps({"products": [{"title": "保温杯", "code": "C-1", "容量": "500ml"}]}, ensure_ascii=False),
                  encoding="utf-8")
    headers, rows = ci.read_table(str(jp))
    products, _ = ci.build_products(rows, ci.suggest_mapping(headers))
    assert products[0]["attributes"] == {"容量": "500ml"}


def test_unsupported_file_type(tmp_path):
    p = tmp_path / "x.pages"
    p.write_text("x")
    with pytest.raises(ValueError):
        ci.read_table(str(p))


def test_import_job_creates_stores_listings_and_reports_progress(svc, tmp_path):
    p = tmp_path / "catalog.csv"
    _write_csv(p, "gbk")
    events = []
    svc.jobs.add_listener(lambda j: events.append((j["status"], j["stage"], j["done"], j["total"])))
    preview = svc.preview_catalog(str(p))
    assert preview["report"]["products"] == 2
    assert preview["stores_in_file"] == ["小铺天猫店", "小铺抖音店"]
    job = svc.import_catalog(str(p), preview["mapping"], run_inline=True)
    job = svc.get_job(job["id"])
    assert job["status"] == "succeeded", job
    assert job["result"]["products"] == 2 and job["progress"] == 1.0
    assert ("running", "index", 2, 2) in events
    stores = {s["name"]: s for s in svc.overview()["stores"]}
    assert stores["小铺天猫店"]["product_count"] == 2
    # The same shirt answers from the douyin store's item id.
    r = svc.search("什么面料", store_ref="小铺抖音店", platform_item_id="4400001")
    assert r["product"]["code"] == "TS-100"

    # Re-importing the same file changes nothing.
    job2 = svc.get_job(svc.import_catalog(str(p), preview["mapping"], run_inline=True)["id"])
    assert job2["result"]["unchanged"] == 2 and job2["result"]["changed"] == 0


def test_cancelled_job(svc, tmp_path):
    p = tmp_path / "catalog.csv"
    _write_csv(p, "gbk")
    mapping = ci.suggest_mapping(HEADERS)
    original = svc.store.upsert_product

    def cancel_then_upsert(*a, **k):
        svc.jobs.cancel(job_holder["id"])
        return original(*a, **k)

    job_holder = {}
    svc.store.upsert_product = cancel_then_upsert
    real_submit = svc.jobs.submit

    def submit(kind, title, fn, tenant_id="", run_inline=False):
        job = svc.store.create_job(kind=kind, title=title, tenant_id=tenant_id)
        job_holder["id"] = job["id"]
        svc.store.update_job(job["id"], status="queued")
        from knowledge.ks.jobs import JobContext, JobCancelled
        ctx = JobContext(svc.jobs, job["id"])
        try:
            fn(ctx)
        except JobCancelled:
            svc.jobs._update(job["id"], status="cancelled", stage="cancelled")
        return svc.store.get_job(job["id"])

    svc.jobs.submit = submit
    job = svc.import_catalog(str(p), mapping)
    svc.jobs.submit = real_submit
    assert svc.get_job(job["id"])["status"] == "cancelled"


def test_docx_product_sheet_draft(svc, tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("商品名称：儿童防水雨衣")
    d.add_paragraph("货号：RC-88")
    d.add_paragraph("面料：涤纶涂层")
    d.add_paragraph("• 反光条夜间更安全")
    d.add_paragraph("Q：可以机洗吗")
    d.add_paragraph("A：建议手洗，不可烘干")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([["尺码", "身高", "衣长"], ["110", "100-110", "60"], ["120", "110-120", "65"]]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    path = tmp_path / "雨衣.docx"
    d.save(path)
    out = svc.draft_from_document(str(path))
    draft = out["draft"]
    assert draft["title"] == "儿童防水雨衣" and draft["code"] == "RC-88"
    assert draft["attributes"]["面料"] == "涤纶涂层"
    assert draft["faq"] == [{"q": "可以机洗吗", "a": "建议手洗，不可烘干"}]
    assert draft["size_chart"]["headers"] == ["尺码", "身高", "衣长"]
    assert draft["selling_points"] == ["反光条夜间更安全"]
    assert out["stats"]["has_size_chart"]

    p = svc.save_product(draft)
    doc = svc.attach_document(p["id"], str(path))
    assert doc["chunk_count"] >= 1
    r = svc.search("身高115穿哪个码", product_id=p["id"])
    assert any(f["key"] == "尺码表" for f in r["facts"])


def test_markdown_table_and_unsupported_doc(tmp_path):
    p = tmp_path / "spec.md"
    p.write_text("# 保温杯\n| 规格 | 值 |\n|---|---|\n| 容量 | 500ml |\n| 材质 | 316不锈钢 |\n", encoding="utf-8")
    text, tables = doc_extract.extract_document(str(p))
    draft = doc_extract.draft_product(str(p), text, tables)
    assert draft["attributes"]["容量"] == "500ml"
    with pytest.raises(ValueError):
        doc_extract.extract_document(str(tmp_path / "a.pptx"))


def test_embedder_change_is_detected_and_reindex_fixes_it(svc, shop):
    assert not svc.overview()["embedder"]["needs_reindex"]
    svc.embedder = HashingEmbedder(dim=256)
    svc.retriever.embedder = svc.embedder
    ov = svc.overview()
    assert ov["embedder"]["needs_reindex"]
    job = svc.get_job(svc.reindex(run_inline=True)["id"])
    assert job["status"] == "succeeded" and job["result"]["chunks"] > 0
    assert not svc.overview()["embedder"]["needs_reindex"]
    r = svc.search("会缩水吗", store_ref=shop["tmall"]["id"], platform_item_id="6600001")
    assert r["faq"] and "1-2%" in r["faq"][0]["answer"]


def test_template_imports_cleanly_with_its_own_mapping(tmp_path):
    path = ci.write_template(str(tmp_path / "模板.xlsx"))
    headers, rows = ci.read_table(path)
    mapping = ci.suggest_mapping(headers)
    assert mapping == {h: f for h, f in ci.TEMPLATE_COLUMNS}  # no manual mapping needed
    products, report = ci.build_products(rows, mapping)
    assert report.errors == [] and report.skipped == [] and report.warnings == []
    by_code = {p["code"]: p for p in products}
    assert set(by_code) == {"DRS-001", "TSH-002"}
    dress = by_code["DRS-001"]
    assert len(dress["variants"]) == 3 and len(dress["faq"]) == 2
    assert dress["selling_points"] == ["收腰显瘦", "透气不闷", "可机洗"]
    assert dress["attributes"]["适用季节"] == "春夏"
    assert all(not p["listings"] for p in products)  # sample rows never create stores

