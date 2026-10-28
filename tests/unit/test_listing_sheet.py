"""商品上架 spreadsheet -> listing.json: store owners fill a spreadsheet, the
上架 skills publish listing.json folders."""
import json
import os

import openpyxl

from agent.ec_skills import listing_sheet


def _sheet(tmp_path, with_main=True):
    p = str(tmp_path / "商品上架.xlsx")
    listing_sheet.write_template(p)
    wb = openpyxl.load_workbook(p)
    wb["商品"].append(["ECT2610B", "珊瑚绒睡衣套装 星月刺绣", "家居服/睡衣/家居服套装", "TONAR", 399, 5,
                     "面料=珊瑚绒；图案=星月\n适用季节=冬季", "加厚保暖"])
    for c, color in enumerate(["燕麦色", "雾霾蓝"], 1):
        for s in ["M", "L", "XL", "XXL"]:
            wb["规格"].append(["ECT2610B", color, s, None, 8 if s == "XXL" else None, f"ECT2610B-{c}{s}"])
    rows = [["ECT2610B", "main_1.png", "主图", None], ["ECT2610B", "detail_1.png", "详情图", None],
            ["ECT2610B", "blue.png", "颜色图", "雾霾蓝"], ["ECT2610B", "chart.png", "尺码表图", None]]
    for r in rows if with_main else rows[1:]:
        wb["图片"].append(r)
    wb["尺码表"].append(["ECT2610B", "M", "160-165", "100-115", 66, 108])
    wb["尺码表"].append(["ECT2610B", "L", "165-170", "115-130", 68, 112])
    wb.save(p)
    for name in ("main_1.png", "detail_1.png", "blue.png", "chart.png"):
        (tmp_path / name).write_bytes(b"x")
    return p


def test_template_has_the_sheets_and_instructions(tmp_path):
    wb = openpyxl.load_workbook(listing_sheet.write_template(str(tmp_path / "t.xlsx")))
    assert wb.sheetnames == ["商品", "规格", "图片", "尺码表", "填写说明"]


def test_a_product_becomes_a_listing_folder(tmp_path):
    out = tmp_path / "inbox"
    [(code, folder, found)] = listing_sheet.convert(_sheet(tmp_path), str(out))
    assert code == "ECT2610B" and found == []
    L = json.load(open(os.path.join(folder, "listing.json"), encoding="utf-8"))
    assert L["category"] == ["家居服", "睡衣", "家居服套装"]
    assert L["attributes"] == {"面料": "珊瑚绒", "图案": "星月", "适用季节": "冬季", "品牌": "TONAR"}
    assert L["colors"] == ["燕麦色", "雾霾蓝"] and L["sizes"] == ["M", "L", "XL", "XXL"]
    assert len(L["skus"]) == 8 and L["skus"][0] == {"颜色": "燕麦色", "尺码": "M", "价格": 399, "库存": 5,
                                                    "商家编码": "ECT2610B-1M"}
    assert L["skus"][3]["库存"] == 8
    assert L["files"] == {"main": ["main_1.png"], "detail": ["detail_1.png"], "color": {"雾霾蓝": "blue.png"},
                          "size_chart": "chart.png"}
    assert L["size_chart"]["columns"] == ["尺码", "身高cm", "体重斤", "衣长cm", "胸围cm"]
    assert L["size_chart"]["rows"][0] == {"尺码": "M", "身高cm": "160-165", "体重斤": "100-115",
                                          "衣长cm": "66", "胸围cm": "108"}
    assert all(os.path.isfile(os.path.join(folder, n)) for n in ("main_1.png", "blue.png", "chart.png"))


def test_problems_are_written_down_for_the_owner(tmp_path):
    [(_, folder, found)] = listing_sheet.convert(_sheet(tmp_path, with_main=False), str(tmp_path / "inbox"))
    assert any("主图" in p for p in found)
    assert json.load(open(os.path.join(folder, "listing.json"), encoding="utf-8"))["problems"] == found


def test_the_spreadsheet_is_not_left_locked(tmp_path):
    # The owner must be able to save or replace it while the app runs (Windows locks open files).
    import shutil
    p = str(tmp_path / "copy.xlsx")
    shutil.copyfile(_sheet(tmp_path), p)        # the helper's own openpyxl handle stays on the original
    listing_sheet.read_sheet(p)
    os.replace(p, p + ".moved")
