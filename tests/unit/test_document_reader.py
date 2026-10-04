"""os_read_document: office / PDF / text documents come back as plain text an LLM can read."""

import json

import pytest

from agent.mcp.server.document_reader import read_document


def test_xlsx_all_sheets_as_rows(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "价格"
    ws.append(["SKU", "Name", "Price"])
    ws.append(["LS-5M", "LED strip 5m", 12.99])
    ws.append(["LB-30", "Light bar", 15.0])
    wb.create_sheet("Notes").append(["ships in 2 days"])
    p = tmp_path / "prices.xlsx"
    wb.save(p)
    out = read_document(str(p))
    assert out["type"] == "xlsx" and out["sheets"] == ["价格", "Notes"]
    assert "LS-5M\tLED strip 5m\t12.99" in out["content"]
    assert "LB-30\tLight bar\t15" in out["content"]
    assert "## Sheet: Notes\nships in 2 days" in out["content"]
    only = read_document(str(p), sheet="Notes")
    assert "LS-5M" not in only["content"]


def test_xlsx_row_cap(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    for i in range(50):
        wb.active.append([i])
    p = tmp_path / "big.xlsx"
    wb.save(p)
    out = read_document(str(p), max_rows=10)
    assert out["truncated"] and "\n9\n" in out["content"] and "\n10\n" not in out["content"]


def test_csv_gbk(tmp_path):
    p = tmp_path / "价格表.csv"
    p.write_bytes("型号,价格\nLS-5M,12.99\n".encode("gbk"))
    out = read_document(str(p))
    assert out["type"] == "csv" and "型号\t价格" in out["content"] and "LS-5M\t12.99" in out["content"]


def test_docx_paragraphs_and_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Lumora 5m RGB LED Strip")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Length", "5 m"
    p = tmp_path / "spec.docx"
    d.save(p)
    out = read_document(str(p))
    assert "Lumora 5m RGB LED Strip" in out["content"] and "Length\t5 m" in out["content"]


def test_pptx_slides(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "16 colours"
    s.notes_slide.notes_text_frame.text = "remote included"
    p = tmp_path / "deck.pptx"
    prs.save(p)
    out = read_document(str(p))
    assert out["slides"] == 1 and "16 colours" in out["content"] and "Notes: remote included" in out["content"]


def test_pdf_text_and_scanned(tmp_path):
    fitz = pytest.importorskip("fitz")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "SKU LS-5M 150 LEDs")
    p = tmp_path / "spec.pdf"
    doc.save(p)
    out = read_document(str(p))
    assert out["pages"] == 1 and "SKU LS-5M 150 LEDs" in out["content"] and "note" not in out
    blank = fitz.open()
    blank.new_page()
    b = tmp_path / "scan.pdf"
    blank.save(b)
    assert "OCR" in read_document(str(b))["note"]


def test_legacy_missing_and_char_cap(tmp_path):
    xls = tmp_path / "old.xls"
    xls.write_bytes(b"\xd0\xcf\x11\xe0")
    assert "save it as .xlsx" in read_document(str(xls))["error"]
    assert read_document(str(tmp_path / "nope.pdf"))["error"] == "file not found"
    t = tmp_path / "long.txt"
    t.write_text("x" * 100, encoding="utf-8")
    out = read_document(str(t), max_chars=10)
    assert out["truncated"] and out["content"].startswith("x" * 10 + "\n...")


def test_tool_wrapper_returns_json(tmp_path):
    import asyncio
    from agent.mcp.server.server import os_read_document
    t = tmp_path / "a.txt"
    t.write_text("hello", encoding="utf-8")
    res = asyncio.run(os_read_document(None, {"input": {"file_path": str(t)}}))
    assert json.loads(res[0].text)["content"] == "hello"
